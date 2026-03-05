import os
import io
from datetime import datetime, date, timedelta
from functools import wraps

from flask import (Flask, render_template, request, jsonify, redirect,
                   url_for, flash, send_file)
from models import db, Site, Category, Criteria, Inspection, InspectionItem

app = Flask(__name__)
app.config['SECRET_KEY'] = 'assessment-platform-secret-key-2026'

import pathlib
_db_dir = pathlib.Path(__file__).resolve().parent / 'data'
_db_dir.mkdir(exist_ok=True)
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{_db_dir}/assessment.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)


def init_demo_data():
    """初始化示例数据，便于快速体验系统"""
    if Category.query.count() > 0:
        return

    categories_data = [
        {'name': '环境卫生', 'description': '现场环境卫生相关考核', 'weight': 1.0, 'sort_order': 1},
        {'name': '安全管理', 'description': '安全生产与安全管理考核', 'weight': 1.5, 'sort_order': 2},
        {'name': '人员管理', 'description': '人员配置与管理考核', 'weight': 1.0, 'sort_order': 3},
        {'name': '设备管理', 'description': '设备维护与管理考核', 'weight': 1.2, 'sort_order': 4},
        {'name': '服务质量', 'description': '服务质量与客户满意度考核', 'weight': 1.3, 'sort_order': 5},
        {'name': '制度执行', 'description': '规章制度执行情况考核', 'weight': 1.0, 'sort_order': 6},
    ]

    for c_data in categories_data:
        cat = Category(**c_data)
        db.session.add(cat)
    db.session.flush()

    criteria_data = [
        {'category': '环境卫生', 'items': [
            ('现场地面清洁', '地面无垃圾、无积水、无油污', 10, '发现一处扣2分'),
            ('垃圾分类与清运', '垃圾按规定分类并及时清运', 10, '未分类扣5分，未及时清运扣3分'),
            ('绿化维护', '绿化区域维护良好，无枯死植物', 5, '发现枯死植物每处扣1分'),
            ('公共区域整洁', '公共区域物品摆放整齐，无杂物堆放', 10, '发现杂物堆放每处扣2分'),
        ]},
        {'category': '安全管理', 'items': [
            ('消防设施完好', '消防器材在有效期内且摆放正确', 15, '过期每个扣5分，摆放不当扣2分'),
            ('安全标识齐全', '危险区域安全警示标识齐全', 10, '缺失一处扣3分'),
            ('安全通道畅通', '安全出口和逃生通道无阻塞', 15, '发现阻塞每处扣5分'),
            ('用电安全', '无私拉乱接电线，电器使用规范', 10, '发现违规用电每处扣5分'),
        ]},
        {'category': '人员管理', 'items': [
            ('人员到岗情况', '按合同要求配足人员', 10, '缺岗一人扣3分'),
            ('着装规范', '工作人员统一着装，佩戴工牌', 5, '未着装扣2分/人，未佩戴工牌扣1分/人'),
            ('培训记录', '定期开展培训并有记录', 5, '无培训记录扣5分'),
            ('考勤管理', '考勤记录完整准确', 5, '记录不完整扣3分'),
        ]},
        {'category': '设备管理', 'items': [
            ('设备运行正常', '主要设备运行状态良好', 15, '设备故障未及时维修每台扣5分'),
            ('设备维保记录', '设备维护保养记录完整', 10, '记录缺失扣5分，不完整扣3分'),
            ('备品备件管理', '常用备品备件储备充足', 5, '储备不足扣3分'),
        ]},
        {'category': '服务质量', 'items': [
            ('服务态度', '工作人员服务态度良好', 10, '态度恶劣每次扣5分'),
            ('响应速度', '问题响应及时，不超过规定时限', 10, '超时一次扣3分'),
            ('投诉处理', '投诉处理及时有效', 10, '未处理投诉每件扣5分'),
        ]},
        {'category': '制度执行', 'items': [
            ('规章制度上墙', '相关规章制度公示到位', 5, '未公示扣5分'),
            ('台账管理', '各类台账记录完整规范', 10, '缺失一项扣3分'),
            ('应急预案', '应急预案完善并定期演练', 10, '无预案扣10分，未演练扣5分'),
        ]},
    ]

    for group in criteria_data:
        cat = Category.query.filter_by(name=group['category']).first()
        for i, (name, desc, max_score, rule) in enumerate(group['items']):
            c = Criteria(
                category_id=cat.id, name=name, description=desc,
                max_score=max_score, deduction_rule=rule, sort_order=i + 1
            )
            db.session.add(c)

    site = Site(name='示例考核地点A', address='XX市XX区XX路XX号',
                contact_person='张三', contact_phone='13800138000',
                description='第一批下属现场考核点')
    db.session.add(site)
    site2 = Site(name='示例考核地点B', address='XX市XX区XX街XX号',
                 contact_person='李四', contact_phone='13900139000',
                 description='第二批下属现场考核点')
    db.session.add(site2)

    db.session.commit()


with app.app_context():
    db.create_all()
    init_demo_data()


# ==================== 页面路由 ====================

@app.route('/')
def dashboard():
    return render_template('dashboard.html')


@app.route('/sites')
def sites_page():
    return render_template('sites.html')


@app.route('/criteria')
def criteria_page():
    return render_template('criteria.html')


@app.route('/inspections')
def inspections_page():
    return render_template('inspections.html')


@app.route('/inspections/new')
def new_inspection_page():
    return render_template('inspection_form.html')


@app.route('/inspections/<int:inspection_id>')
def inspection_detail_page(inspection_id):
    return render_template('inspection_detail.html', inspection_id=inspection_id)


@app.route('/inspections/<int:inspection_id>/edit')
def edit_inspection_page(inspection_id):
    return render_template('inspection_form.html', inspection_id=inspection_id)


# ==================== API: 仪表盘统计 ====================

@app.route('/api/dashboard/stats')
def api_dashboard_stats():
    total_sites = Site.query.count()
    total_inspections = Inspection.query.count()
    total_criteria = Criteria.query.filter_by(is_active=True).count()
    completed_inspections = Inspection.query.filter_by(status='completed').count()

    month_start = date.today().replace(day=1)
    month_inspections = Inspection.query.filter(
        Inspection.inspection_date >= month_start
    ).count()

    avg_score = db.session.query(db.func.avg(Inspection.total_score)).filter(
        Inspection.status.in_(['completed', 'reviewed'])
    ).scalar() or 0

    pending_issues = InspectionItem.query.filter_by(
        rectification_status='pending'
    ).count()

    return jsonify({
        'total_sites': total_sites,
        'total_inspections': total_inspections,
        'total_criteria': total_criteria,
        'completed_inspections': completed_inspections,
        'month_inspections': month_inspections,
        'avg_score': round(float(avg_score), 1),
        'pending_issues': pending_issues,
    })


@app.route('/api/dashboard/trend')
def api_dashboard_trend():
    months = int(request.args.get('months', 6))
    today = date.today()
    result = []

    for i in range(months - 1, -1, -1):
        d = today.replace(day=1) - timedelta(days=i * 30)
        month_start = d.replace(day=1)
        if d.month == 12:
            month_end = d.replace(year=d.year + 1, month=1, day=1)
        else:
            month_end = d.replace(month=d.month + 1, day=1)

        inspections = Inspection.query.filter(
            Inspection.inspection_date >= month_start,
            Inspection.inspection_date < month_end,
            Inspection.status.in_(['completed', 'reviewed'])
        ).all()

        avg = sum(i.total_score for i in inspections) / len(inspections) if inspections else 0
        result.append({
            'month': month_start.strftime('%Y-%m'),
            'count': len(inspections),
            'avg_score': round(avg, 1),
        })

    return jsonify(result)


@app.route('/api/dashboard/category_scores')
def api_dashboard_category_scores():
    categories = Category.query.order_by(Category.sort_order).all()
    result = []

    for cat in categories:
        criteria_ids = [c.id for c in cat.criteria.all()]
        if not criteria_ids:
            continue

        items = InspectionItem.query.filter(
            InspectionItem.criteria_id.in_(criteria_ids)
        ).all()

        if items:
            total_max = sum(item.max_score for item in items if hasattr(item, 'max_score'))
            total_deduction = sum(item.deduction for item in items)
            qualified = sum(1 for item in items if item.is_qualified)
            result.append({
                'category': cat.name,
                'total_deduction': round(total_deduction, 1),
                'qualified_rate': round(qualified / len(items) * 100, 1) if items else 100,
                'check_count': len(items),
            })
        else:
            result.append({
                'category': cat.name,
                'total_deduction': 0,
                'qualified_rate': 100,
                'check_count': 0,
            })

    return jsonify(result)


@app.route('/api/dashboard/site_rankings')
def api_dashboard_site_rankings():
    sites = Site.query.all()
    result = []
    for site in sites:
        inspections = site.inspections.filter(
            Inspection.status.in_(['completed', 'reviewed'])
        ).all()
        if inspections:
            avg = sum(i.total_score for i in inspections) / len(inspections)
            latest = max(inspections, key=lambda i: i.inspection_date)
            result.append({
                'site_name': site.name,
                'avg_score': round(avg, 1),
                'latest_score': latest.total_score,
                'inspection_count': len(inspections),
            })
    result.sort(key=lambda x: x['avg_score'], reverse=True)
    return jsonify(result)


@app.route('/api/dashboard/recent_issues')
def api_dashboard_recent_issues():
    items = InspectionItem.query.filter(
        InspectionItem.is_qualified == False
    ).join(Inspection).order_by(
        Inspection.inspection_date.desc()
    ).limit(10).all()

    result = []
    for item in items:
        result.append({
            'site_name': item.inspection.site.name if item.inspection and item.inspection.site else '',
            'criteria_name': item.criteria_ref.name if item.criteria_ref else '',
            'deduction': item.deduction,
            'issue': item.issue_description or '',
            'date': item.inspection.inspection_date.strftime('%Y-%m-%d') if item.inspection else '',
            'rectification_status': item.rectification_status,
        })
    return jsonify(result)


# ==================== API: 考核地点管理 ====================

@app.route('/api/sites', methods=['GET'])
def api_get_sites():
    sites = Site.query.order_by(Site.created_at.desc()).all()
    return jsonify([s.to_dict() for s in sites])


@app.route('/api/sites', methods=['POST'])
def api_create_site():
    data = request.get_json()
    site = Site(
        name=data['name'],
        address=data.get('address', ''),
        contact_person=data.get('contact_person', ''),
        contact_phone=data.get('contact_phone', ''),
        description=data.get('description', ''),
    )
    db.session.add(site)
    db.session.commit()
    return jsonify(site.to_dict()), 201


@app.route('/api/sites/<int:site_id>', methods=['GET'])
def api_get_site(site_id):
    site = Site.query.get_or_404(site_id)
    return jsonify(site.to_dict())


@app.route('/api/sites/<int:site_id>', methods=['PUT'])
def api_update_site(site_id):
    site = Site.query.get_or_404(site_id)
    data = request.get_json()
    site.name = data.get('name', site.name)
    site.address = data.get('address', site.address)
    site.contact_person = data.get('contact_person', site.contact_person)
    site.contact_phone = data.get('contact_phone', site.contact_phone)
    site.description = data.get('description', site.description)
    db.session.commit()
    return jsonify(site.to_dict())


@app.route('/api/sites/<int:site_id>', methods=['DELETE'])
def api_delete_site(site_id):
    site = Site.query.get_or_404(site_id)
    db.session.delete(site)
    db.session.commit()
    return jsonify({'message': '删除成功'})


# ==================== API: 考核类别管理 ====================

@app.route('/api/categories', methods=['GET'])
def api_get_categories():
    cats = Category.query.order_by(Category.sort_order).all()
    return jsonify([c.to_dict() for c in cats])


@app.route('/api/categories', methods=['POST'])
def api_create_category():
    data = request.get_json()
    cat = Category(
        name=data['name'],
        description=data.get('description', ''),
        weight=data.get('weight', 1.0),
        sort_order=data.get('sort_order', 0),
    )
    db.session.add(cat)
    db.session.commit()
    return jsonify(cat.to_dict()), 201


@app.route('/api/categories/<int:cat_id>', methods=['PUT'])
def api_update_category(cat_id):
    cat = Category.query.get_or_404(cat_id)
    data = request.get_json()
    cat.name = data.get('name', cat.name)
    cat.description = data.get('description', cat.description)
    cat.weight = data.get('weight', cat.weight)
    cat.sort_order = data.get('sort_order', cat.sort_order)
    db.session.commit()
    return jsonify(cat.to_dict())


@app.route('/api/categories/<int:cat_id>', methods=['DELETE'])
def api_delete_category(cat_id):
    cat = Category.query.get_or_404(cat_id)
    db.session.delete(cat)
    db.session.commit()
    return jsonify({'message': '删除成功'})


# ==================== API: 考核标准管理 ====================

@app.route('/api/criteria', methods=['GET'])
def api_get_criteria():
    category_id = request.args.get('category_id')
    query = Criteria.query
    if category_id:
        query = query.filter_by(category_id=int(category_id))
    criteria_list = query.order_by(Criteria.category_id, Criteria.sort_order).all()
    return jsonify([c.to_dict() for c in criteria_list])


@app.route('/api/criteria', methods=['POST'])
def api_create_criteria():
    data = request.get_json()
    c = Criteria(
        category_id=data['category_id'],
        name=data['name'],
        description=data.get('description', ''),
        max_score=data.get('max_score', 100),
        deduction_rule=data.get('deduction_rule', ''),
        sort_order=data.get('sort_order', 0),
    )
    db.session.add(c)
    db.session.commit()
    return jsonify(c.to_dict()), 201


@app.route('/api/criteria/<int:criteria_id>', methods=['PUT'])
def api_update_criteria(criteria_id):
    c = Criteria.query.get_or_404(criteria_id)
    data = request.get_json()
    c.category_id = data.get('category_id', c.category_id)
    c.name = data.get('name', c.name)
    c.description = data.get('description', c.description)
    c.max_score = data.get('max_score', c.max_score)
    c.deduction_rule = data.get('deduction_rule', c.deduction_rule)
    c.sort_order = data.get('sort_order', c.sort_order)
    c.is_active = data.get('is_active', c.is_active)
    db.session.commit()
    return jsonify(c.to_dict())


@app.route('/api/criteria/<int:criteria_id>', methods=['DELETE'])
def api_delete_criteria(criteria_id):
    c = Criteria.query.get_or_404(criteria_id)
    db.session.delete(c)
    db.session.commit()
    return jsonify({'message': '删除成功'})


# ==================== API: 考核记录管理 ====================

@app.route('/api/inspections', methods=['GET'])
def api_get_inspections():
    site_id = request.args.get('site_id')
    status = request.args.get('status')
    query = Inspection.query
    if site_id:
        query = query.filter_by(site_id=int(site_id))
    if status:
        query = query.filter_by(status=status)
    inspections = query.order_by(Inspection.inspection_date.desc()).all()
    return jsonify([i.to_dict() for i in inspections])


@app.route('/api/inspections', methods=['POST'])
def api_create_inspection():
    data = request.get_json()
    inspection = Inspection(
        site_id=data['site_id'],
        inspector=data['inspector'],
        inspection_date=datetime.strptime(data['inspection_date'], '%Y-%m-%d').date(),
        notes=data.get('notes', ''),
        status='draft',
    )
    db.session.add(inspection)
    db.session.flush()

    all_criteria = Criteria.query.filter_by(is_active=True).order_by(
        Criteria.category_id, Criteria.sort_order
    ).all()
    for c in all_criteria:
        item = InspectionItem(
            inspection_id=inspection.id,
            criteria_id=c.id,
            score=c.max_score,
            deduction=0,
            is_qualified=True,
        )
        db.session.add(item)

    inspection.total_score = sum(c.max_score for c in all_criteria)
    inspection.total_deduction = 0
    db.session.commit()
    return jsonify(inspection.to_dict()), 201


@app.route('/api/inspections/<int:inspection_id>', methods=['GET'])
def api_get_inspection(inspection_id):
    inspection = Inspection.query.get_or_404(inspection_id)
    result = inspection.to_dict()
    result['items'] = [item.to_dict() for item in inspection.items.join(Criteria).order_by(
        Criteria.category_id, Criteria.sort_order
    ).all()]
    return jsonify(result)


@app.route('/api/inspections/<int:inspection_id>', methods=['PUT'])
def api_update_inspection(inspection_id):
    inspection = Inspection.query.get_or_404(inspection_id)
    data = request.get_json()

    if 'inspector' in data:
        inspection.inspector = data['inspector']
    if 'inspection_date' in data:
        inspection.inspection_date = datetime.strptime(data['inspection_date'], '%Y-%m-%d').date()
    if 'notes' in data:
        inspection.notes = data['notes']
    if 'status' in data:
        inspection.status = data['status']

    db.session.commit()
    return jsonify(inspection.to_dict())


@app.route('/api/inspections/<int:inspection_id>', methods=['DELETE'])
def api_delete_inspection(inspection_id):
    inspection = Inspection.query.get_or_404(inspection_id)
    db.session.delete(inspection)
    db.session.commit()
    return jsonify({'message': '删除成功'})


# ==================== API: 考核明细管理 ====================

@app.route('/api/inspection-items/<int:item_id>', methods=['PUT'])
def api_update_inspection_item(item_id):
    item = InspectionItem.query.get_or_404(item_id)
    data = request.get_json()

    if 'deduction' in data:
        item.deduction = float(data['deduction'])
        criteria = Criteria.query.get(item.criteria_id)
        item.score = max(0, criteria.max_score - item.deduction)
        item.is_qualified = item.deduction == 0

    if 'issue_description' in data:
        item.issue_description = data['issue_description']
    if 'rectification_status' in data:
        item.rectification_status = data['rectification_status']
    if 'notes' in data:
        item.notes = data['notes']

    inspection = Inspection.query.get(item.inspection_id)
    all_items = inspection.items.all()
    inspection.total_deduction = sum(i.deduction for i in all_items)
    total_max = sum(
        Criteria.query.get(i.criteria_id).max_score for i in all_items
    )
    inspection.total_score = total_max - inspection.total_deduction

    db.session.commit()

    result = item.to_dict()
    result['inspection_total_score'] = inspection.total_score
    result['inspection_total_deduction'] = inspection.total_deduction
    return jsonify(result)


@app.route('/api/inspection-items/batch', methods=['PUT'])
def api_batch_update_items():
    data = request.get_json()
    items_data = data.get('items', [])
    inspection_id = None

    for item_data in items_data:
        item = InspectionItem.query.get(item_data['id'])
        if not item:
            continue
        inspection_id = item.inspection_id

        if 'deduction' in item_data:
            item.deduction = float(item_data['deduction'])
            criteria = Criteria.query.get(item.criteria_id)
            item.score = max(0, criteria.max_score - item.deduction)
            item.is_qualified = item.deduction == 0
        if 'issue_description' in item_data:
            item.issue_description = item_data['issue_description']
        if 'rectification_status' in item_data:
            item.rectification_status = item_data['rectification_status']
        if 'notes' in item_data:
            item.notes = item_data['notes']

    if inspection_id:
        inspection = Inspection.query.get(inspection_id)
        all_items = inspection.items.all()
        inspection.total_deduction = sum(i.deduction for i in all_items)
        total_max = sum(
            Criteria.query.get(i.criteria_id).max_score for i in all_items
        )
        inspection.total_score = total_max - inspection.total_deduction

    db.session.commit()
    return jsonify({'message': '批量更新成功'})


# ==================== API: 数据导出 ====================

@app.route('/api/export/inspection/<int:inspection_id>')
def api_export_inspection(inspection_id):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return jsonify({'error': '导出功能需要安装 openpyxl 库'}), 500

    inspection = Inspection.query.get_or_404(inspection_id)
    items = inspection.items.join(Criteria).order_by(
        Criteria.category_id, Criteria.sort_order
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '考核记录'

    header_font = Font(bold=True, size=14)
    sub_font = Font(size=11)
    table_header_font = Font(bold=True, size=10, color='FFFFFF')
    table_header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.merge_cells('A1:H1')
    ws['A1'] = '现场考核记录表'
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align

    ws['A3'] = '考核地点：'
    ws['B3'] = inspection.site.name
    ws['D3'] = '检查人员：'
    ws['E3'] = inspection.inspector
    ws['A4'] = '考核日期：'
    ws['B4'] = inspection.inspection_date.strftime('%Y-%m-%d')
    ws['D4'] = '总得分：'
    ws['E4'] = inspection.total_score
    ws['A5'] = '状态：'
    ws['B5'] = inspection.to_dict()['status_label']
    ws['D5'] = '总扣分：'
    ws['E5'] = inspection.total_deduction

    if inspection.notes:
        ws['A6'] = '备注：'
        ws['B6'] = inspection.notes

    headers = ['序号', '考核类别', '考核项目', '考核标准', '满分', '扣分', '得分', '问题描述', '整改状态']
    row = 8
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = table_header_font
        cell.fill = table_header_fill
        cell.alignment = center_align
        cell.border = border

    for idx, item in enumerate(items, 1):
        row += 1
        values = [
            idx,
            item.criteria_ref.category.name if item.criteria_ref and item.criteria_ref.category else '',
            item.criteria_ref.name if item.criteria_ref else '',
            item.criteria_ref.description if item.criteria_ref else '',
            item.criteria_ref.max_score if item.criteria_ref else 0,
            item.deduction,
            item.score,
            item.issue_description or '',
            item.to_dict()['rectification_label'],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 8
    ws.column_dimensions['F'].width = 8
    ws.column_dimensions['G'].width = 8
    ws.column_dimensions['H'].width = 30
    ws.column_dimensions['I'].width = 12

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f'考核记录_{inspection.site.name}_{inspection.inspection_date.strftime("%Y%m%d")}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@app.route('/api/export/summary')
def api_export_summary():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return jsonify({'error': '导出功能需要安装 openpyxl 库'}), 500

    inspections = Inspection.query.filter(
        Inspection.status.in_(['completed', 'reviewed'])
    ).order_by(Inspection.inspection_date.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '考核汇总'

    header_font = Font(bold=True, size=14)
    table_header_font = Font(bold=True, size=10, color='FFFFFF')
    table_header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.merge_cells('A1:G1')
    ws['A1'] = '现场考核汇总表'
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align
    ws['A2'] = f'导出时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}'

    headers = ['序号', '考核地点', '考核日期', '检查人员', '总得分', '总扣分', '状态']
    row = 4
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = table_header_font
        cell.fill = table_header_fill
        cell.alignment = center_align
        cell.border = border

    for idx, insp in enumerate(inspections, 1):
        row += 1
        values = [
            idx,
            insp.site.name if insp.site else '',
            insp.inspection_date.strftime('%Y-%m-%d'),
            insp.inspector,
            insp.total_score,
            insp.total_deduction,
            insp.to_dict()['status_label'],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'考核汇总_{date.today().strftime("%Y%m%d")}.xlsx',
    )


if __name__ == '__main__':
    import sys
    debug = '--debug' in sys.argv
    app.run(debug=debug, host='0.0.0.0', port=5000, use_reloader=False)
